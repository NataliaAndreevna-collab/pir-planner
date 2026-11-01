"""
Генератор коммерческого предложения (.docx) на фирменном бланке организации ООО «ТПС».
"""
import io
import os
from datetime import date as date_cls
from datetime import datetime
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.models import (
    KPPaymentScheme,
    KPPaymentStage,
    KPPriceFormat,
    KPRequest,
    KPStage,
    KPTermUnit,
)
from app.services.money import format_rub, plural_form, rubles_in_words

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DEFAULT_LETTERHEAD = os.path.join(BASE_DIR, "data", "private", "kp_letterhead.docx")

MONTHS_GEN = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря"
]

STAGE_TEXT = {
    KPStage.PD: "разработку проектной документации",
    KPStage.RD: "разработку рабочей документации",
    KPStage.PD_RD: "разработку проектной документации (стадия ПД + РД)",
    KPStage.CHANGES: "внесение изменений в проектную документацию",
}

DEFAULT_PAYMENT_STAGES = [
    KPPaymentStage(
        percent=40,
        text="Предоплата - 40% стоимости работ по договору производится в течении 5-ти рабочих дней с момента заключения договора;"
    ),
    KPPaymentStage(
        percent=30,
        text="Оплата в размере 30% стоимости работ производится в течении 5 рабочих дней с момента передачи полного комплекта проектной документации заказчику;"
    ),
    KPPaymentStage(
        percent=30,
        text="Оплата в размере 30% стоимости работ производится в течении 5 рабочих дней с момента подписания акта выполненных работ."
    ),
]


def get_letterhead_path() -> Optional[str]:
    path = os.environ.get("KP_LETTERHEAD_PATH") or DEFAULT_LETTERHEAD
    return path if os.path.isfile(path) else None


# ---------------------------------------------------------------- вспомогательное

def _parse_date(value: str) -> date_cls:
    if value:
        try:
            return datetime.strptime(value[:10], "%Y-%m-%d").date()
        except ValueError:
            pass
    return date_cls.today()


def format_kp_date(d: date_cls) -> str:
    return f"«{d.day:02d}» {MONTHS_GEN[d.month - 1]} {d.year}г."


def outgoing_number_text(number: str, d: date_cls) -> str:
    suffix = f"{d.month:02d}{d.year % 100:02d}"
    clean = number.strip() or "____"
    return f"Исх.№{clean}/{suffix}"


def term_text(value: int, unit: KPTermUnit) -> str:
    forms = {
        KPTermUnit.MONTHS: ("месяц", "месяца", "месяцев"),
        KPTermUnit.WORKING_DAYS: ("рабочий день", "рабочих дня", "рабочих дней"),
        KPTermUnit.CALENDAR_DAYS: ("календарный день", "календарных дня", "календарных дней"),
    }[unit]
    return f"{value} {plural_form(value, forms)}"


def _vat_info(req: KPRequest):
    cp = req.project.cost_parameters
    if cp and cp.include_vat:
        return True, cp.vat_rate
    return False, 0.0


def _rate_str(rate: float) -> str:
    return f"{rate:g}%"


def _format_cost(val: float, with_kopecks: bool = True) -> str:
    if with_kopecks:
        return f"{val:,.2f}".replace(",", " ").replace(".", ",")
    return f"{int(round(val)):,}".replace(",", " ")


def _rows_and_totals(req: KPRequest):
    """Возвращает (строки таблицы, итого без НДС, НДС, итого с НДС)."""
    cp = req.project.cost_parameters
    include_vat, rate = _vat_info(req)
    rows = []
    if req.price_format == KPPriceFormat.SECTIONS and cp:
        for i, it in enumerate([s for s in cp.sections_costs if s.enabled], 1):
            rows.append((str(i), it.code_mark, it.title, it.cost_rub))
        for aw in getattr(cp, "additional_works", []) or []:
            if aw.enabled and aw.cost_net_rub > 0:
                rows.append((str(len(rows) + 1), "Доп.", aw.name, aw.cost_net_rub))
        net = round(sum(r[3] for r in rows), 2)
    elif req.price_format == KPPriceFormat.WORKS:
        for i, w in enumerate(req.works, 1):
            rows.append((str(i), "", w.name, w.cost_rub))
        net = round(sum(r[3] for r in rows), 2)
    else:
        net = round(cp.total_cost_net_rub, 2) if cp else 0.0
    vat = round(net * rate / 100.0, 2) if include_vat else 0.0
    return rows, net, vat, round(net + vat, 2)


# ---------------------------------------------------------------- построение текста

def _set_font(run, size=None, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _para(
    doc,
    text="",
    bold=False,
    italic=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    size=11,
    space_after=6,
    space_before=0,
    first_line_indent=None,
    right_indent=None,
    left_indent=None,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if first_line_indent is not None:
        p.paragraph_format.first_line_indent = first_line_indent
    if right_indent is not None:
        p.paragraph_format.right_indent = right_indent
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
    if text:
        _set_font(p.add_run(text), size, bold, italic)
    return p


def _cell_text(cell, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    _set_font(p.add_run(text), size, bold, italic)


def _set_borders(table):
    """Границы таблицы вручную."""
    from docx.oxml import OxmlElement
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def _price_table(doc, req: KPRequest, rows, net, vat, gross):
    include_vat, rate = _vat_info(req)
    is_works = (req.price_format == KPPriceFormat.WORKS)

    if is_works:
        # Формат по видам работ (3 колонки: № п/п, Наименование работ, Стоимость, руб.)
        table = doc.add_table(rows=1, cols=3)
        _set_borders(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Cm(1.2), Cm(12.8), Cm(4.0)]

        hdr = table.rows[0].cells
        for c, t in zip(hdr, ["№ п/п", "Наименование работ", "Стоимость, руб."]):
            _cell_text(c, t, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)

        for num, _, name, cost in rows:
            r = table.add_row().cells
            _cell_text(r[0], num, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
            _cell_text(r[1], name, align=WD_ALIGN_PARAGRAPH.LEFT, size=11)
            _cell_text(r[2], _format_cost(cost, with_kopecks=True), align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)

        # Строка ИТОГО:
        r_itogo = table.add_row().cells
        _cell_text(r_itogo[0], "", size=11)
        _cell_text(r_itogo[1], "ИТОГО:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=11)
        _cell_text(r_itogo[2], _format_cost(net, with_kopecks=True), bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)

    else:
        # Формат по разделам ПД/РД (4 колонки: № п/п, Марка, Наименование, Стоимость, руб.)
        table = doc.add_table(rows=1, cols=4)
        _set_borders(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Cm(1.2), Cm(1.8), Cm(11.0), Cm(4.0)]

        hdr = table.rows[0].cells
        for c, t in zip(hdr, ["№ п/п", "Марка", "Наименование", "Стоимость, руб."]):
            _cell_text(c, t, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)

        stage_header = "Проектная и рабочая документация" if req.stage == KPStage.PD_RD else (
            "Рабочая документация" if req.stage == KPStage.RD else "Проектная документация"
        )
        r_stage = table.add_row().cells
        merged_stage = r_stage[0].merge(r_stage[1]).merge(r_stage[2]).merge(r_stage[3])
        _cell_text(merged_stage, stage_header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)

        for num, mark, name, cost in rows:
            r = table.add_row().cells
            _cell_text(r[0], num, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
            _cell_text(r[1], mark, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
            _cell_text(r[2], name, align=WD_ALIGN_PARAGRAPH.LEFT, size=11)
            _cell_text(r[3], _format_cost(cost, with_kopecks=True), align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)

        r_itogo = table.add_row().cells
        merged_itogo = r_itogo[0].merge(r_itogo[1]).merge(r_itogo[2])
        _cell_text(merged_itogo, "ИТОГО:", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)
        _cell_text(r_itogo[3], _format_cost(net, with_kopecks=True), bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)

        if include_vat:
            r_vat = table.add_row().cells
            merged_vat = r_vat[0].merge(r_vat[1]).merge(r_vat[2])
            _cell_text(merged_vat, f"ИТОГО, с учетом НДС {_rate_str(rate)}:", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)
            _cell_text(r_vat[3], _format_cost(gross, with_kopecks=True), bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)

    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    return table


def _build_body(doc, req: KPRequest):
    """Основной текст КП: обращение, предмет, таблица, стоимость, оговорки, срок, оплата."""
    include_vat, rate = _vat_info(req)
    rows, net, vat, gross = _rows_and_totals(req)

    # 1. Обращение: размер 11 pt, Полужирный Курсив, по центру
    _para(doc, req.salutation, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=4)

    # Дополнительная пустая строка после обращения
    _para(doc, "", size=11, space_after=4)

    # 2. Предмет КП: размер 11 pt, по ширине с красной строкой (1.25 см)
    if req.subject_text.strip():
        subject = req.subject_text.strip()
    else:
        ref = f"В ответ на Ваш запрос от {req.request_date}, " if req.request_date else "По Вашему запросу "
        subject = (f"{ref}направляем коммерческое предложение на {STAGE_TEXT[req.stage]} "
                   f"по объекту: «{req.project.object_name}»:")
    _para(doc, subject, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, first_line_indent=Cm(1.25))

    # 3. Входит в стоимость (если указано)
    if req.inclusions:
        _para(doc, "В коммерческое предложение входит:", size=11, space_after=2, first_line_indent=Cm(1.25))
        for item in req.inclusions:
            _para(doc, "- " + item.rstrip(";."), size=11, space_after=2, left_indent=Cm(1.25))
        _para(doc, "", size=11, space_after=4)

    # 4. Таблица стоимости
    if req.price_format != KPPriceFormat.TOTAL and rows:
        _price_table(doc, req, rows, net, vat, gross)
        # Пустая строка после таблицы убрана!

    # 5. Строка стоимости сразу под таблицей (space_before=0, красная строка 1.25 см)
    amount = gross if include_vat else net
    vat_note = f", в т.ч. НДС {_rate_str(rate)}" if include_vat else ", НДС не облагается"
    _para(
        doc,
        f"Стоимость работ составляет {rubles_in_words(amount)}{vat_note}.",
        bold=True,
        italic=True,
        size=11,
        space_after=6,
        space_before=0,
        first_line_indent=Cm(1.25),
    )

    # 6. Исключения (с красной строкой 1.25 см)
    if req.exclusions:
        _para(
            doc,
            "*Стоимость не включает в себя: " + "; ".join(e.rstrip(";.") for e in req.exclusions) + ".",
            size=11,
            space_after=6,
            first_line_indent=Cm(1.25),
        )

    # 7. Условия оплаты (с красной строкой 1.25 см)
    _para(doc, "Условия оплаты:", bold=True, size=11, space_after=2, first_line_indent=Cm(1.25))
    if req.payment_scheme == KPPaymentScheme.STAGES or req.payment_stages:
        stages = req.payment_stages or DEFAULT_PAYMENT_STAGES
        if len(stages) == 1:
            _para(doc, stages[0].text.strip(), size=11, space_after=4, first_line_indent=Cm(1.25))
        else:
            _para(doc, "Оплата работ производится поэтапно, с авансированием:", size=11, space_after=2, first_line_indent=Cm(1.25))
            for st in stages:
                _para(doc, "- " + st.text.strip().lstrip("- "), size=11, space_after=2, left_indent=Cm(1.25))
    else:
        _para(doc, f"Оплата производится на условиях авансирования - {req.advance_percent:g}%.", size=11, space_after=4, first_line_indent=Cm(1.25))

    _para(doc, "", size=11, space_after=2)

    # 8. Срок выполнения (красная строка)
    if req.term_value:
        _para(doc, f"Срок выполнения работ – {term_text(req.term_value, req.term_unit)}.", size=11, space_after=4, first_line_indent=Cm(1.25))

    # 9. Оговорка по согласованию (красная строка)
    if req.note_approval_excluded:
        _para(doc, "*В срок не входит время согласования документации с заказчиком и заинтересованными организациями.", size=11, space_after=4, first_line_indent=Cm(1.25))

    # 10. Готовы помочь в разработке ТЗ (красная строка)
    if req.phrase_help_tz:
        _para(doc, "* Готовы помочь в разработке Технического задания", italic=True, size=11, space_after=4, first_line_indent=Cm(1.25))

    # 11. Встречные предложения (красная строка)
    if req.phrase_counter_offers:
        _para(doc, "* Готовы рассмотреть встречные предложения по срокам и стоимости проектирования.",
              bold=True, italic=True, size=11, space_after=4, first_line_indent=Cm(1.25))

    # 12. Действительность КП (красная строка)
    _para(doc, f"Коммерческое предложение действительно - {req.validity_days} рабочих дней.", size=11, space_after=6, first_line_indent=Cm(1.25))


# ---------------------------------------------------------------- оформление подписи и бланка

def _format_signer_paragraph(p, req: KPRequest):
    """
    Форматирует строку подписи директора:
    'Директор ООО «ТПС» ___________________________________/{ФИО}'
    """
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = req.signer_title.strip() or "Директор"
    company = req.company_name.strip() or "ООО «ТПС»"
    name = (req.signer_name.strip() or "Валова М.М.").lstrip("/")

    r1 = p.add_run(f"{title} {company} ")
    _set_font(r1, size=11, bold=True, italic=True)

    # Удлиненное нижнее подчеркивание (35 символов)
    r2 = p.add_run("___________________________________")
    _set_font(r2, size=11, bold=False, italic=False)

    r3 = p.add_run(f"/{name} ")
    _set_font(r3, size=11, bold=True, italic=True)


def _format_executor_paragraph(p, req: KPRequest):
    """
    Форматирует строку исполнителя строго по образцу ТПС:
    'Исп. Эберт Е.В. Тел. +7 932 326 93 83'
    """
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    raw = req.executor.strip() or "Эберт Е.В. Тел. +7 932 326 93 83"
    exec_str = raw if raw.startswith("Исп.") else f"Исп. {raw}"

    r = p.add_run(exec_str)
    _set_font(r, size=10, bold=False, italic=True)


def _build_corporate(req: KPRequest, d: date_cls, letterhead: str) -> Document:
    doc = Document(letterhead)
    mapping = {
        "{{OUT_NUM}}": outgoing_number_text(req.outgoing_number, d),
        "{{DATE}}": format_kp_date(d),
        "{{ADDRESSEE}}": "\n".join(req.addressee_lines),
    }
    anchor = None
    for p in doc.paragraphs:
        if "{{BODY}}" in p.text:
            anchor = p
        elif "{{SIGNER}}" in p.text:
            _format_signer_paragraph(p, req)
        elif "{{EXECUTOR}}" in p.text:
            _format_executor_paragraph(p, req)
        else:
            for key, val in mapping.items():
                if key in p.text:
                    p.text = p.text.replace(key, val)

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "{{BODY}}" in p.text:
                        anchor = p
                    elif "{{SIGNER}}" in p.text:
                        _format_signer_paragraph(p, req)
                    elif "{{EXECUTOR}}" in p.text:
                        _format_executor_paragraph(p, req)
                    else:
                        for key, val in mapping.items():
                            if key in p.text:
                                p.text = p.text.replace(key, val)

    # Две пустые строки перед «По объекту: ...» в левом блоке шапки
    if len(doc.tables) > 1 and len(doc.tables[1].rows) > 0:
        c0 = doc.tables[1].rows[0].cells[0]
        while len(c0.paragraphs) < 6:
            c0.add_paragraph()
        c0.paragraphs[4].text = ""
        c0.paragraphs[4].paragraph_format.space_before = Pt(0)
        c0.paragraphs[4].paragraph_format.space_after = Pt(0)
        c0.paragraphs[5].text = ""
        c0.paragraphs[5].paragraph_format.space_before = Pt(0)
        c0.paragraphs[5].paragraph_format.space_after = Pt(0)

        p_obj = c0.paragraphs[6] if len(c0.paragraphs) > 6 else c0.add_paragraph()
        p_obj.text = ""
        r = p_obj.add_run(f"По объекту: «{req.project.object_name}»")
        _set_font(r, size=10, bold=True, italic=True)
        p_obj.paragraph_format.space_before = Pt(0)
        p_obj.paragraph_format.space_after = Pt(0)

    if anchor is None:
        raise ValueError("В бланке не найдена метка {{BODY}}")

    body = doc.element.body
    before = list(body)
    _build_body(doc, req)
    new_elements = [el for el in body if el not in before and el.tag != qn("w:sectPr")]
    for el in new_elements:
        anchor._element.addprevious(el)
    anchor._element.getparent().remove(anchor._element)
    return doc


def generate_kp_docx(req: KPRequest) -> io.BytesIO:
    d = _parse_date(req.date)
    letterhead = get_letterhead_path()
    if not letterhead:
        raise FileNotFoundError(
            "Фирменный бланк компании не найден (ожидается файл data/private/kp_letterhead.docx или KP_LETTERHEAD_PATH). "
            "Формирование КП доступно только на фирменном бланке организации."
        )
    doc = _build_corporate(req, d, letterhead)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
