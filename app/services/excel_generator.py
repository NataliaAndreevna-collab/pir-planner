import calendar
import datetime
import io
import math
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.models import ProjectData, ProjectMode
from app.services.pp87_rules import (
    GROUPS, GROUP_COLORS, resolve_schedule, parse_date, format_date_ru,
    total_calendar_days, total_weeks,
)

MONTH_NAMES_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"
]

GROUP_EXCEL_COLORS = {
    "precontract": "475569", # Тёмно-серый / сланец (Договор)
    "survey": "10B981",      # Изумрудный (Изыскания)
    "tx_concept": "F97316",  # Оранжевый (Технология)
    "pd": "2563EB",          # Синий (ПД)
    "approval": "D97706",    # Янтарный (Согласования)
    "internal_rvk": "0284C7",# Сине-голубой / морской (Внутренняя экспертиза РВК)
    "expertise": "9333EA",   # Пурпурный (Экспертиза)
    "rd": "0D9488",          # Бирюзовый (РД)
    "other": "64748B",       # Серо-голубой (Прочее / РнС)
}


def generate_excel_schedule(project_data: ProjectData) -> io.BytesIO:
    """
    Генерирует профессиональный Календарный график выполнения работ в Excel (.xlsx)
    с двухуровневой шкалой (Месяцы + Декады/Числа 1-10, 11-20, 21-30/31), связями этапов, группами и вехами.
    """
    resolve_schedule(project_data.schedule_phases, project_data.start_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Календарный график"
    ws.views.sheetView[0].showGridLines = True

    steel_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sub_hdr_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    week_hdr_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    group_row_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    milestone_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    font_title = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=9.5)
    font_small = Font(name="Calibri", size=8.5, italic=True)
    font_white_bold = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
    font_week = Font(name="Calibri", size=8, bold=True, color="1E3A8A")
    font_bar = Font(name="Calibri", size=8.5, bold=True, color="FFFFFF")
    font_milestone = Font(name="Calibri", size=10, bold=True, color="B45309")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Базовая дата старта проекта
    base_dt = parse_date(project_data.start_date) or datetime.date.today()
    base_y = base_dt.year
    base_m = base_dt.month

    # Определяем горизонт проекта в месяцах и декадах (числах 1-10, 11-20, 21-конец)
    enabled_phases = [p for p in project_data.schedule_phases if getattr(p, "enabled", True)]
    max_finish_dt = base_dt
    for p in enabled_phases:
        f_dt = parse_date(p.finish_date)
        if f_dt and f_dt > max_finish_dt:
            max_finish_dt = f_dt

    # Общее число месяцев (минимум 2 месяца)
    month_diff = (max_finish_dt.year - base_y) * 12 + (max_finish_dt.month - base_m) + 1
    num_months = max(2, month_diff)
    num_decades = num_months * 3

    first_decade_col = 14
    note_col = first_decade_col + num_decades

    current_row = 1

    # 1. Шапка документа
    if project_data.mode == ProjectMode.CONTRACT:
        ws.cell(row=current_row, column=note_col, value="Приложение № 2").font = font_small
        ws.cell(
            row=current_row + 1,
            column=note_col,
            value=f"к договору № {project_data.contract_number or '____'} от {project_data.contract_date or '___'}"
        ).font = font_small
        current_row += 3

        ws.cell(row=current_row, column=2, value="ИСПОЛНИТЕЛЬ:").font = font_bold
        ws.cell(row=current_row, column=first_decade_col, value="ЗАКАЗЧИК:").font = font_bold

        contr_name = project_data.contractor or project_data.customer or "Заказчик"
        ws.cell(row=current_row + 1, column=2, value="________________ / ___________ /").font = font_regular
        ws.cell(row=current_row + 1, column=first_decade_col, value=f"{contr_name}").font = font_regular
        ws.cell(row=current_row + 2, column=first_decade_col, value="________________ / ___________ /").font = font_regular
        current_row += 4
    else:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=note_col)
        c_title = ws.cell(row=current_row, column=1, value="КАЛЕНДАРНЫЙ ГРАФИК ВЫПОЛНЕНИЯ ПРОЕКТНО-ИЗЫСКАТЕЛЬСКИХ РАБОТ")
        c_title.font = font_title
        c_title.alignment = align_center
        current_row += 2

    # Объект и сводка сроков
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=note_col)
    actual_days = total_calendar_days(project_data.schedule_phases, project_data.start_date)
    actual_months = round(actual_days / 30.4375, 1)
    summary_str = f"Объект: {project_data.object_name}   |   Срок: {actual_days} кал. дн. (~{actual_months} мес.)   |   Период: {format_date_ru(base_dt)} – {format_date_ru(max_finish_dt)}"
    if project_data.customer:
        summary_str += f"   |   Заказчик: {project_data.customer}"
    if project_data.contractor:
        summary_str += f"   |   Генпроектировщик: {project_data.contractor}"
    c_obj = ws.cell(row=current_row, column=1, value=summary_str)
    c_obj.font = Font(name="Calibri", size=10.5, bold=True, color="1E3A8A")
    c_obj.alignment = align_left
    current_row += 2

    # 2. Двухуровневая шапка таблицы
    tbl_start_row = current_row
    sub_row = tbl_start_row + 1

    fixed_headers = [
        (1, "№\nэтапа"),
        (2, "Наименование этапа / раздела"),
        (3, "Ответственный"),
        (4, "Статус"),
        (5, "Начало"),
        (6, "Длительность"),
        (7, "Окончание"),
        (8, "Передано\nЗаказчику"),
        (9, "Замечания 1"),
        (10, "Ответы 1"),
        (11, "Замечания 2"),
        (12, "Ответы 2"),
        (13, "Результат\nсогласования"),
    ]
    for col_idx, text in fixed_headers:
        ws.merge_cells(start_row=tbl_start_row, start_column=col_idx, end_row=sub_row, end_column=col_idx)
        c = ws.cell(row=tbl_start_row, column=col_idx, value=text)
        c.fill = steel_fill
        c.font = font_white_bold
        c.alignment = align_center

    # Шкала: Месяцы (по 3 колонки) + Числа месяца (1-10, 11-20, 21-30/31)
    for m_idx in range(num_months):
        cur_m = (base_m - 1 + m_idx) % 12 + 1
        cur_y = base_y + (base_m - 1 + m_idx) // 12
        _, last_day = calendar.monthrange(cur_y, cur_m)

        d_start_col = first_decade_col + m_idx * 3
        d_end_col = d_start_col + 2
        ws.merge_cells(start_row=tbl_start_row, start_column=d_start_col, end_row=tbl_start_row, end_column=d_end_col)

        m_label = f"{MONTH_NAMES_RU[cur_m - 1]} {cur_y}"
        cm = ws.cell(row=tbl_start_row, column=d_start_col, value=m_label)
        cm.fill = sub_hdr_fill
        cm.font = Font(name="Calibri", size=9.5, bold=True, color="1E3A8A")
        cm.alignment = align_center

        # Подшапка: точные числа месяца
        decade_labels = ["1-10", "11-20", f"21-{last_day}"]
        for d_i, d_lbl in enumerate(decade_labels):
            col_idx = d_start_col + d_i
            cw = ws.cell(row=sub_row, column=col_idx, value=d_lbl)
            cw.fill = week_hdr_fill
            cw.font = font_week
            cw.alignment = align_center

    # Колонка Примечание
    ws.merge_cells(start_row=tbl_start_row, start_column=note_col, end_row=sub_row, end_column=note_col)
    c_note = ws.cell(row=tbl_start_row, column=note_col, value="Примечание")
    c_note.fill = steel_fill
    c_note.font = font_white_bold
    c_note.alignment = align_center

    for r in (tbl_start_row, sub_row):
        for col in range(1, note_col + 1):
            ws.cell(row=r, column=col).border = thin_border

    current_row = sub_row + 1

    # 3. Вывод строк графика с группировкой по блокам
    last_group = None
    for idx, phase in enumerate(enabled_phases, start=1):
        grp = phase.group or "pd"
        if grp != last_group:
            last_group = grp
            grp_title = GROUPS.get(grp, "Этапы работ")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=note_col)
            cg = ws.cell(row=current_row, column=1, value=grp_title.upper())
            cg.fill = group_row_fill
            cg.font = Font(name="Calibri", size=9.5, bold=True, color="1E3A8A")
            cg.alignment = align_left
            for col in range(1, note_col + 1):
                ws.cell(row=current_row, column=col).border = thin_border
            current_row += 1

        p_start_dt = parse_date(phase.start_date) or base_dt
        p_finish_dt = parse_date(phase.finish_date) or p_start_dt
        dur_days = 0 if phase.milestone else max(1, getattr(phase, "duration_days", 14))

        start_str = format_date_ru(p_start_dt)
        finish_str = format_date_ru(p_finish_dt)
        dur_str = "Веха" if phase.milestone else f"{dur_days} дн."

        row_vals = [
            (1, phase.id or idx, align_center, font_bold if phase.milestone else font_regular),
            (2, f"◆ {phase.name}" if phase.milestone else phase.name, align_left, font_bold if phase.milestone else font_regular),
            (3, getattr(phase, "resp_party", "Проектировщик"), align_center, font_regular),
            (4, getattr(phase, "status", "В работе"), align_center, font_regular),
            (5, start_str, align_center, font_regular),
            (6, dur_str, align_center, font_regular),
            (7, finish_str, align_center, font_regular),
            (8, getattr(phase, "sent_to_client_date", "") or "—", align_center, font_regular),
            (9, getattr(phase, "client_remarks_1_date", "") or "—", align_center, font_regular),
            (10, getattr(phase, "designer_answers_1_date", "") or "—", align_center, font_regular),
            (11, getattr(phase, "client_remarks_2_date", "") or "—", align_center, font_regular),
            (12, getattr(phase, "designer_answers_2_date", "") or "—", align_center, font_regular),
            (13, getattr(phase, "approval_result", "") or "—", align_left, font_regular),
        ]
        for col_idx, val, al, fnt in row_vals:
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = al
            cell.font = fnt
            cell.border = thin_border
            if phase.milestone:
                cell.fill = milestone_fill

        # Индексы декад для Ганта
        start_dec_idx = (p_start_dt.year - base_y) * 12 + (p_start_dt.month - base_m)
        start_dec_idx = start_dec_idx * 3 + (0 if p_start_dt.day <= 10 else (1 if p_start_dt.day <= 20 else 2))

        finish_dec_idx = (p_finish_dt.year - base_y) * 12 + (p_finish_dt.month - base_m)
        finish_dec_idx = finish_dec_idx * 3 + (0 if p_finish_dt.day <= 10 else (1 if p_finish_dt.day <= 20 else 2))

        bar_hex = GROUP_EXCEL_COLORS.get(grp, "3B82F6")
        bar_fill = PatternFill(start_color=bar_hex, end_color=bar_hex, fill_type="solid")

        for d_col in range(num_decades):
            col_idx = first_decade_col + d_col
            cw = ws.cell(row=current_row, column=col_idx)
            cw.border = thin_border
            if phase.milestone:
                cw.fill = milestone_fill
                if d_col == start_dec_idx:
                    cw.value = "◆"
                    cw.alignment = align_center
                    cw.font = font_milestone
            else:
                if start_dec_idx <= d_col <= finish_dec_idx:
                    cw.fill = bar_fill
                    if d_col == start_dec_idx:
                        cw.value = f"{dur_days}д"
                        cw.alignment = align_center
                        cw.font = font_bar

        c_nt = ws.cell(row=current_row, column=note_col, value=phase.note)
        c_nt.alignment = align_left
        c_nt.font = font_small
        c_nt.border = thin_border
        if phase.milestone:
            c_nt.fill = milestone_fill

        current_row += 1

    # Настройка ширины колонок Листа 1
    ws.column_dimensions["A"].width = 7.0
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 12
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 12
    ws.column_dimensions["L"].width = 12
    ws.column_dimensions["M"].width = 20
    for d_col in range(num_decades):
        col_letter = get_column_letter(first_decade_col + d_col)
        ws.column_dimensions[col_letter].width = 6.0
    ws.column_dimensions[get_column_letter(note_col)].width = 32

    ws.row_dimensions[tbl_start_row].height = 22
    ws.row_dimensions[sub_row].height = 20

    # =========================================================================
    # ЛИСТ 2: РЕЕСТР СОГЛАСОВАНИЙ И ЗАМЕЧАНИЙ (РАБОЧИЙ ЖУРНАЛ ГИПА)
    # По образцу листа "Геолог"
    # =========================================================================
    ws2 = wb.create_sheet(title="Реестр замечаний и ответов")
    ws2.views.sheetView[0].showGridLines = True

    # Заголовок Листа 2
    ws2.merge_cells("A1:K1")
    title2 = ws2.cell(row=1, column=1, value="РЕЕСТР СОГЛАСОВАНИЙ, ЗАМЕЧАНИЙ И ОТВЕТОВ (ЖУРНАЛ ГИПА)")
    title2.font = font_title
    title2.alignment = align_center

    ws2.merge_cells("A2:K2")
    sub_title2 = ws2.cell(row=2, column=1, value=f"Объект: {project_data.object_name}   |   Заказчик: {project_data.customer or 'Не указан'}")
    sub_title2.font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
    sub_title2.alignment = align_left

    reg_headers = [
        (1, "№\nп/п", 6),
        (2, "Вид работы / Раздел документации", 40),
        (3, "Статус\nсогласования", 16),
        (4, "Плановый\nсрок", 12),
        (5, "Дата направления\nЗаказчику / в эксп.", 18),
        (6, "Исходящий №\nпроектировщика", 18),
        (7, "Замечания Заказчика / Экспертизы\n(вх. №, дата, суть замечаний)", 35),
        (8, "Ответы проектировщика\n(исх. №, дата, статус устранения)", 35),
        (9, "Повторные замечания\n(вх. №, дата)", 25),
        (10, "Финальный ответ / Снятие\n(исх. №, протокол)", 25),
        (11, "Итоговый результат\n/ Заключение", 25),
    ]

    r2_hdr = 4
    for col_idx, h_text, w_width in reg_headers:
        c2 = ws2.cell(row=r2_hdr, column=col_idx, value=h_text)
        c2.fill = steel_fill
        c2.font = font_white_bold
        c2.alignment = align_center
        c2.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = w_width

    ws2.row_dimensions[r2_hdr].height = 28

    r2_cur = 5
    # Заполняем строками из ключевых этапов (согласования, экспертиза, тома ПД/РД)
    target_phases = [p for p in enabled_phases if p.group in ("approval", "expertise", "tx_concept", "precontract") or not p.milestone]
    for p_idx, phase in enumerate(target_phases, start=1):
        status_val = getattr(phase, "status", "В работе")
        c_status_fill = None
        if "Согласовано" in status_val or "Завершено" in status_val:
            c_status_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Зеленоватый
        elif "Замечания" in status_val or "На согласовании" in status_val:
            c_status_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Желтоватый

        row2_data = [
            (1, p_idx, align_center, font_bold if phase.milestone else font_regular),
            (2, phase.name, align_left, font_bold if phase.milestone else font_regular),
            (3, status_val, align_center, font_bold if "Согласовано" in status_val else font_regular),
            (4, "Веха" if phase.milestone else f"{phase.duration_weeks} нед.", align_center, font_regular),
            (5, getattr(phase, "sent_to_client_date", "") or "", align_center, font_regular),
            (6, "", align_center, font_regular), # Место для ввода исх. номера
            (7, getattr(phase, "client_remarks_1_date", "") or "", align_left, font_regular),
            (8, getattr(phase, "designer_answers_1_date", "") or "", align_left, font_regular),
            (9, getattr(phase, "client_remarks_2_date", "") or "", align_left, font_regular),
            (10, getattr(phase, "designer_answers_2_date", "") or "", align_left, font_regular),
            (11, getattr(phase, "approval_result", "") or ("Согласовано" if phase.status == "Согласовано" else "В процессе"), align_left, font_regular),
        ]

        for col_idx, val, al, fnt in row2_data:
            cell = ws2.cell(row=r2_cur, column=col_idx, value=val)
            cell.alignment = al
            cell.font = fnt
            cell.border = thin_border
            if col_idx == 3 and c_status_fill:
                cell.fill = c_status_fill
            elif phase.milestone:
                cell.fill = milestone_fill

        ws2.row_dimensions[r2_cur].height = 20
        r2_cur += 1

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
