import io
from typing import List
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from app.models import FacilityType, ProjectData, ProjectMode, SectionItem


FONT_FAMILY = "Times New Roman"


def _set_run_font(run, size_pt: float = 11, bold: bool = False, italic: bool = False, color_rgb: RGBColor = None):
    run.font.name = FONT_FAMILY
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FAMILY)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb is not None:
        run.font.color.rgb = color_rgb


def set_cell_background(cell, fill_hex: str):
    """Устанавливает цвет фона ячейки."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_table_borders(table, outer_color="1E293B", inner_color="94A3B8", outer_sz="6", inner_sz="4"):
    """Устанавливает строгие проектные рамки таблицы по ГОСТ."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{outer_sz}" w:space="0" w:color="{outer_color}"/>'
        f'  <w:left w:val="single" w:sz="{outer_sz}" w:space="0" w:color="{outer_color}"/>'
        f'  <w:bottom w:val="single" w:sz="{outer_sz}" w:space="0" w:color="{outer_color}"/>'
        f'  <w:right w:val="single" w:sz="{outer_sz}" w:space="0" w:color="{outer_color}"/>'
        f'  <w:insideH w:val="single" w:sz="{inner_sz}" w:space="0" w:color="{inner_color}"/>'
        f'  <w:insideV w:val="single" w:sz="{inner_sz}" w:space="0" w:color="{inner_color}"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


def set_cell_margins(cell, top=90, bottom=90, left=130, right=130):
    """Устанавливает внутренние отступы ячейки (в dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def _repeat_table_header(row):
    """Повторять строку шапки таблицы при переносе на следующую страницу."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tr_pr.append(tbl_header)


def _prevent_row_split(row):
    """Запретить разрыв строки таблицы между страницами."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _add_page_number_field(run):
    """Добавляет поле номера страницы Word (PAGE)."""
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld_char1)
    r.append(instr_text)
    r.append(fld_char2)
    r.append(fld_char3)


def _setup_footer(section, project_data: ProjectData):
    """Оформляет нижний колонтитул в проектном стиле (шифр + номер листа)."""
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cipher_label = f"Шифр: {project_data.cipher}   |   " if project_data.cipher else ""
    r1 = p.add_run(f"{cipher_label}Ведомость состава документации   |   Лист ")
    _set_run_font(r1, size_pt=9, italic=True, color_rgb=RGBColor(100, 116, 139))
    r_pg = p.add_run()
    _set_run_font(r_pg, size_pt=9, bold=True, color_rgb=RGBColor(71, 85, 105))
    _add_page_number_field(r_pg)


def _render_passport_table(doc, project_data: ProjectData, pd_count: int, rd_count: int):
    """Отрисовывает паспортную карточку объекта проектирования."""
    rows_data = [
        ("Наименование объекта проектирования:", project_data.object_name or "—"),
    ]
    if project_data.customer:
        rows_data.append(("Заказчик (Застройщик):", project_data.customer))
    if project_data.contractor:
        rows_data.append(("Генеральный проектировщик:", project_data.contractor))

    fac_label = (
        "Линейный объект (Раздел III ПП РФ № 87)"
        if project_data.facility_type == FacilityType.LINEAR
        else "Объект капитального строительства производственного / непроизводственного назначения (Раздел II ПП РФ № 87)"
    )
    rows_data.append(("Вид объекта строительства:", fac_label))

    stages_list = []
    if pd_count > 0:
        stages_list.append(f"Проектная документация (Стадия «П» — {pd_count} том.)")
    if rd_count > 0:
        stages_list.append(f"Рабочая документация (Стадия «Р» — {rd_count} компл.)")
    rows_data.append(("Стадийность проектирования:", "; ".join(stages_list) if stages_list else "Проектная документация"))

    if project_data.cipher:
        rows_data.append(("Базовый шифр документации:", project_data.cipher))

    preset_map = {
        "tz": "В соответствии с перечнем Технического задания",
        "permit": "Комплект для получения Разрешения на строительство (ч. 7 ст. 51 ГрК РФ) с учетом ТЗ",
        "expertise": "Комплект для прохождения Экспертизы проектной документации (ПП РФ № 87, ГОСТ Р 21.101-2020)",
    }
    norm_text = preset_map.get(getattr(project_data, "composition_preset", "permit"), preset_map["permit"])
    rows_data.append(("Нормативная основа формирования состава:", norm_text))

    table = doc.add_table(rows=len(rows_data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, outer_color="334155", inner_color="CBD5E1", outer_sz="6", inner_sz="4")

    col_widths = [Cm(6.0), Cm(12.0)]
    for row_idx, (label, val) in enumerate(rows_data):
        row = table.rows[row_idx]
        _prevent_row_split(row)
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(c0, "F1F5F9")
        set_cell_margins(c0, top=75, bottom=75, left=120, right=120)
        set_cell_margins(c1, top=75, bottom=75, left=120, right=120)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        _set_run_font(r0, size_pt=10, bold=True, color_rgb=RGBColor(30, 41, 59))

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        is_main = row_idx == 0 or label.startswith("Базовый шифр")
        r1 = p1.add_run(val)
        _set_run_font(r1, size_pt=10.5 if row_idx == 0 else 10, bold=is_main)


def _render_sections_table(
    doc,
    sections: List[SectionItem],
    vol_col_title: str,
    code_col_title: str,
    name_col_title: str,
    header_fill_hex: str,
    summary_label: str,
):
    """Отрисовывает ведомость разделов ПД или основных комплектов РД по ГОСТ Р 21.101-2020."""
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, outer_color="1E293B", inner_color="94A3B8", outer_sz="6", inner_sz="4")

    # Итого ширина: 1.2 + 2.3 + 3.5 + 8.0 + 3.0 = 18.0 см (лист А4 с полями 2.0 см слева и 1.0 см справа)
    col_widths = [Cm(1.2), Cm(2.3), Cm(3.5), Cm(8.0), Cm(3.0)]

    hdr_row = table.rows[0]
    _repeat_table_header(hdr_row)
    _prevent_row_split(hdr_row)

    headers = ["№ п/п", vol_col_title, code_col_title, name_col_title, "Примечание"]
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, header_fill_hex)
        set_cell_margins(cell, top=110, bottom=110, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        _set_run_font(run, size_pt=10, bold=True, color_rgb=RGBColor(255, 255, 255))

    last_header = None
    for idx_sec, sec in enumerate(sections, start=1):
        if sec.category_header and sec.category_header != last_header:
            last_header = sec.category_header
            cat_row = table.add_row()
            _prevent_row_split(cat_row)
            merged_cell = cat_row.cells[0]
            for c in cat_row.cells[1:]:
                merged_cell.merge(c)
            set_cell_background(merged_cell, "E2E8F0")
            set_cell_margins(merged_cell, top=85, bottom=85, left=130, right=130)
            p_cat = merged_cell.paragraphs[0]
            p_cat.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_cat.paragraph_format.space_before = Pt(0)
            p_cat.paragraph_format.space_after = Pt(0)
            r_cat = p_cat.add_run(sec.category_header)
            _set_run_font(r_cat, size_pt=10, bold=True, color_rgb=RGBColor(30, 58, 138))

        row = table.add_row()
        _prevent_row_split(row)
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = w
            set_cell_margins(row.cells[idx], top=75, bottom=75, left=100, right=100)
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # 0: № п/п
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(str(idx_sec))
        _set_run_font(r0, size_pt=10)

        # 1: Номер тома / книги
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        vol_str = str(sec.volume_num or idx_sec)
        if not vol_str.lower().startswith(("том", "книга", "рд")):
            vol_str = f"Том {vol_str}" if getattr(sec, "stage", "П") != "Р" else f"Книга {idx_sec}"
        r1 = p1.add_run(vol_str)
        _set_run_font(r1, size_pt=10, bold=True)

        # 2: Обозначение (шифр)
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(sec.full_code or sec.code_mark)
        _set_run_font(r2, size_pt=10, bold=True, color_rgb=RGBColor(30, 58, 138))

        # 3: Наименование
        p3 = row.cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(sec.title)
        _set_run_font(r3, size_pt=10)

        # 4: Примечание
        p4 = row.cells[4].paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p4.paragraph_format.space_before = Pt(0)
        p4.paragraph_format.space_after = Pt(0)
        r4 = p4.add_run(sec.note or "")
        _set_run_font(r4, size_pt=9.5, italic=True, color_rgb=RGBColor(71, 85, 105))

    # Итоговая строка по таблице
    sum_row = table.add_row()
    _prevent_row_split(sum_row)
    merged_sum = sum_row.cells[0]
    for c in sum_row.cells[1:]:
        merged_sum.merge(c)
    set_cell_background(merged_sum, "F8FAFC")
    set_cell_margins(merged_sum, top=85, bottom=85, left=130, right=130)
    p_sum = merged_sum.paragraphs[0]
    p_sum.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sum.paragraph_format.space_before = Pt(0)
    p_sum.paragraph_format.space_after = Pt(0)
    r_sum = p_sum.add_run(summary_label)
    _set_run_font(r_sum, size_pt=10, bold=True, color_rgb=RGBColor(30, 41, 59))


def generate_docx_sostav(project_data: ProjectData) -> io.BytesIO:
    """
    Генерирует документ Word «Ведомость состава проектной и рабочей документации»
    в соответствии со стандартами оформления проектной документации (ГОСТ Р 21.101-2020, ПП РФ № 87).
    """
    doc = Document()

    # Поля листа по ГОСТ Р 21.101-2020: левое 20 мм (под брошюровку), остальные по 10–15 мм
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.0)
        _setup_footer(section, project_data)

    pd_sections = [s for s in project_data.sections if s.enabled and getattr(s, "stage", "П") != "Р"]
    rd_sections = [s for s in project_data.sections if s.enabled and getattr(s, "stage", "П") == "Р"]

    # 1. Гриф Приложения и блок подписей (в режиме Договора)
    if project_data.mode == ProjectMode.CONTRACT:
        p_app = doc.add_paragraph()
        p_app.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_app.paragraph_format.space_after = Pt(6)
        r_app = p_app.add_run(
            f"Приложение № 1\nк Договору № {project_data.contract_number or '_______'} "
            f"от {project_data.contract_date or '«____» ___________ 2026 г.'}"
        )
        _set_run_font(r_app, size_pt=10, italic=True)

        table_sig = doc.add_table(rows=1, cols=2)
        table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_sig.autofit = False

        left_cell, right_cell = table_sig.rows[0].cells[0], table_sig.rows[0].cells[1]
        left_cell.width, right_cell.width = Cm(9.0), Cm(9.0)

        contr_name = project_data.contractor or project_data.customer or "Заказчик"

        p_left = left_cell.paragraphs[0]
        p_left.paragraph_format.space_after = Pt(6)
        r_l1 = p_left.add_run("ОТ ИСПОЛНИТЕЛЯ:\n")
        _set_run_font(r_l1, size_pt=10, bold=True)
        r_l2 = p_left.add_run("Руководитель проектной организации\n\n________________ / ________________ /\n«____» ________________ 202__ г.\nМ.П.")
        _set_run_font(r_l2, size_pt=10)

        p_right = right_cell.paragraphs[0]
        p_right.paragraph_format.space_after = Pt(6)
        r_r1 = p_right.add_run("ОТ ЗАКАЗЧИКА:\n")
        _set_run_font(r_r1, size_pt=10, bold=True)
        r_r2 = p_right.add_run(f"{contr_name}\n\n________________ / ________________ /\n«____» ________________ 202__ г.\nМ.П.")
        _set_run_font(r_r2, size_pt=10)

    # 2. Заголовок Ведомости
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(8)

    if pd_sections and rd_sections:
        header_title = "ВЕДОМОСТЬ СОСТАВА ПРОЕКТНОЙ И РАБОЧЕЙ ДОКУМЕНТАЦИИ"
    elif rd_sections:
        header_title = "ВЕДОМОСТЬ ОСНОВНЫХ КОМПЛЕКТОВ РАБОЧИХ ЧЕРТЕЖЕЙ (РАБОЧАЯ ДОКУМЕНТАЦИЯ)"
    else:
        header_title = "ВЕДОМОСТЬ СОСТАВА ПРОЕКТНОЙ ДОКУМЕНТАЦИИ"

    r_title = p_title.add_run(header_title)
    _set_run_font(r_title, size_pt=13, bold=True, color_rgb=RGBColor(30, 58, 138))

    # 3. Паспорт объекта проектирования
    _render_passport_table(doc, project_data, len(pd_sections), len(rd_sections))

    # 4. Таблица Стадии «П» (Проектная документация по ПП РФ № 87)
    if pd_sections:
        p_pd = doc.add_paragraph()
        p_pd.paragraph_format.space_before = Pt(12)
        p_pd.paragraph_format.space_after = Pt(4)
        pd_heading = (
            "1. Состав проектной документации (Стадия «П» — в соответствии с Постановлением Правительства РФ от 16.02.2008 № 87)"
            if rd_sections
            else "Состав проектной документации (Стадия «П» — в соответствии с Постановлением Правительства РФ от 16.02.2008 № 87)"
        )
        r_pd = p_pd.add_run(pd_heading)
        _set_run_font(r_pd, size_pt=11, bold=True, color_rgb=RGBColor(30, 41, 59))

        _render_sections_table(
            doc,
            pd_sections,
            vol_col_title="Номер тома",
            code_col_title="Обозначение (шифр раздела)",
            name_col_title="Наименование раздела / подраздела проектной документации",
            header_fill_hex="1E3A8A",
            summary_label=f"Итого по Стадии «П» (Проектная документация): {len(pd_sections)} томов (книг)",
        )

    # 5. Таблица Стадии «Р» (Рабочая документация по ГОСТ Р 21.101-2020)
    if rd_sections:
        p_rd = doc.add_paragraph()
        p_rd.paragraph_format.space_before = Pt(12)
        p_rd.paragraph_format.space_after = Pt(4)
        rd_heading = (
            "2. Ведомость основных комплектов рабочих чертежей (Стадия «Р» — в соответствии с ГОСТ Р 21.101-2020)"
            if pd_sections
            else "Ведомость основных комплектов рабочих чертежей (Стадия «Р» — в соответствии с ГОСТ Р 21.101-2020)"
        )
        r_rd = p_rd.add_run(rd_heading)
        _set_run_font(r_rd, size_pt=11, bold=True, color_rgb=RGBColor(15, 118, 110))

        _render_sections_table(
            doc,
            rd_sections,
            vol_col_title="Номер книги",
            code_col_title="Обозначение (марка РД)",
            name_col_title="Наименование основного комплекта рабочих чертежей",
            header_fill_hex="0F766E",
            summary_label=f"Итого по Стадии «Р» (Рабочая документация): {len(rd_sections)} основных комплектов",
        )

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
